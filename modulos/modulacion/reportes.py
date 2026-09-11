"""
Reporte "Programa de despacho": un .xlsx con las modulaciones de una fecha,
agrupadas y ordenadas por el alias del cliente, con el formato de la hoja de
control de la terminal.

Módulo puro: construye el Workbook; no toca request/response.
"""
import datetime
import re
from itertools import groupby

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from django.utils import timezone

from .models import Modulacion

_MESES_ES = {
    1: 'ene', 2: 'feb', 3: 'mar', 4: 'abr', 5: 'may', 6: 'jun',
    7: 'jul', 8: 'ago', 9: 'sep', 10: 'oct', 11: 'nov', 12: 'dic',
}

ENCABEZADOS = [
    'FECHA DE DESPACHO', 'TERMINAL DE DESPACHO', 'AGENCIA', 'TIPO', 'PESO',
    'CONTENEDOR', 'OPERADOR', 'CLIENTE', 'CARRIL',
    'REGISTRO', '', '',
    'SELLO', 'Nº DE MANIOBRA',
]
_ANCHOS = [12, 16, 10, 6, 8, 16, 34, 12, 8, 14, 14, 14, 7, 12]

_FILL_HEADER = PatternFill('solid', fgColor='D9D9D9')
_FILL_MORADO = PatternFill('solid', fgColor='E4DFEC')
_FILL_GRUPO = PatternFill('solid', fgColor='F2F2F2')
_FILL_REG = {
    10: PatternFill('solid', fgColor='C6EFCE'),   # J verde
    11: PatternFill('solid', fgColor='E4DFEC'),   # K morado
    12: PatternFill('solid', fgColor='DDEBF7'),   # L azul
}
_THIN = Side(border_style='thin', color='BFBFBF')
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
_CENTRO = Alignment(horizontal='center', vertical='center')
_CENTRO_WRAP = Alignment(horizontal='center', vertical='center', wrap_text=True)
_WRAP = Alignment(wrap_text=True, vertical='center')


def _fecha_es(d):
    return f'{d.day:02d}-{_MESES_ES[d.month]}-{d:%y}'


def _hora_es(dt):
    h = dt.hour % 12 or 12
    sufijo = 'a. m.' if dt.hour < 12 else 'p. m.'
    return f'{h}:{dt.minute:02d} {sufijo}'


def _texto_hora(dt, fecha_reporte):
    if dt is None:
        return ''
    local = timezone.localtime(dt) if timezone.is_aware(dt) else dt
    txt = _hora_es(local)
    if local.date() == fecha_reporte:
        txt += ' (HOY)'
    return txt


def _linea_unidad(unidad):
    if unidad is None:
        return ''
    return f'ECO {unidad.numero_economico} PLACAS {unidad.placa}'


def _tipo(m):
    # Solo los dígitos iniciales: '40HC' -> '40', '45G1' -> '45'.
    m_ini = re.match(r'\d+', m.tipo_contenedor or '')
    return m_ini.group(0) if m_ini else (m.tipo_contenedor or '')


_SIN_CLIENTE = '—'


def _etiqueta_grupo(m):
    return m.cliente.etiqueta if m.cliente_id else _SIN_CLIENTE


def _clave_grupo(m):
    """Clave de agrupación/orden por alias, normalizada (mayús/minús)."""
    return _etiqueta_grupo(m).casefold()


def _clave_orden(m):
    term = m.terminal_portuaria.etiqueta if m.terminal_portuaria_id else ''
    return (
        m.cliente_id is None,   # las modulaciones sin cliente van al final
        _clave_grupo(m),
        term.casefold(),
        m.hora_registro is None,
        m.hora_registro or datetime.datetime.min,
    )


def construir_programa_despacho(fecha):
    modulaciones = sorted(
        Modulacion.objects
        .filter(fecha_modulacion_aduana=fecha)
        .select_related('cliente', 'operador', 'unidad', 'agencia', 'terminal_portuaria'),
        key=_clave_orden,
    )

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Programa de despacho'

    for col, texto in enumerate(ENCABEZADOS, start=1):
        c = ws.cell(row=1, column=col, value=texto)
        c.font = Font(bold=True)
        c.fill = _FILL_HEADER
        c.border = _BORDER
        c.alignment = _CENTRO_WRAP
    ws.merge_cells('J1:L1')
    ws['J1'].value = 'REGISTRO'
    ws['G1'].fill = _FILL_MORADO
    ws['H1'].fill = _FILL_MORADO
    for i, ancho in enumerate(_ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = 'A2'

    fila = 2
    maniobra = 0
    for _clave, grupo in groupby(modulaciones, key=_clave_grupo):
        grupo = list(grupo)
        etiqueta = _etiqueta_grupo(grupo[0])
        ws.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=14)
        gc = ws.cell(row=fila, column=1,
                     value=f'{etiqueta} — {len(grupo)} maniobra(s)')
        gc.font = Font(bold=True)
        gc.fill = _FILL_GRUPO
        fila += 1

        for m in grupo:
            maniobra += 1
            horas = [m.hora_registro, m.hora_ingreso, m.hora_carga]
            presentes = [h for h in horas if h is not None]
            operador_txt = (
                f'{m.operador.nombre}\n{_linea_unidad(m.unidad)}'.rstrip()
                if m.operador_id else _linea_unidad(m.unidad)
            )
            valores = [
                _fecha_es(fecha),
                m.terminal_portuaria.etiqueta if m.terminal_portuaria_id else '',
                m.agencia.nombre if m.agencia_id else '',
                _tipo(m),
                float(m.peso_toneladas) if m.peso_toneladas is not None else None,
                m.contenedor,
                operador_txt,
                _etiqueta_grupo(m),
                m.carril or 'NA',
                None, None, None,
                'SÍ' if m.sello_colocado else 'NO',
                maniobra,
            ]
            for col, v in enumerate(valores, start=1):
                c = ws.cell(row=fila, column=col, value=v)
                c.border = _BORDER
            ws.cell(row=fila, column=5).number_format = '0.00'
            ws.cell(row=fila, column=7).alignment = _WRAP
            ws.cell(row=fila, column=8).font = Font(bold=True)
            ws.row_dimensions[fila].height = 30

            if len(presentes) >= 2:
                for col, h in ((10, horas[0]), (11, horas[1]), (12, horas[2])):
                    c = ws.cell(row=fila, column=col, value=_texto_hora(h, fecha))
                    c.border = _BORDER
                    c.alignment = _CENTRO
                    if h is not None:
                        c.fill = _FILL_REG[col]
            else:
                ws.merge_cells(start_row=fila, start_column=10, end_row=fila, end_column=12)
                texto = _texto_hora(presentes[0], fecha) if presentes else 'CITA PENDIENTE'
                c = ws.cell(row=fila, column=10, value=texto)
                c.alignment = _CENTRO
                for col in (10, 11, 12):
                    ws.cell(row=fila, column=col).border = _BORDER
            fila += 1

    return wb
