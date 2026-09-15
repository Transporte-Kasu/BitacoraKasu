import datetime
from decimal import Decimal
from io import BytesIO

import openpyxl
from django.contrib.auth import get_user_model
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from modulos.bitacoras.models import Cliente
from modulos.modulacion.models import Agencia, Modulacion, TerminalPortuaria
from modulos.modulacion.reportes import construir_programa_despacho
from modulos.operadores.models import Operador
from modulos.unidades.models import Unidad

FECHA = datetime.date(2026, 8, 28)


def _aware(y, mo, d, h, mi):
    return timezone.make_aware(datetime.datetime(y, mo, d, h, mi))


def _mod(**kw):
    kw.setdefault('agencia', Agencia.objects.get_or_create(nombre='LOGINCO')[0])
    kw.setdefault(
        'terminal_portuaria',
        TerminalPortuaria.objects.get_or_create(
            nombre='L.C. Terminal', defaults={'nombre_corto': 'LCTPC'})[0],
    )
    kw.setdefault('tipo_contenedor', '40HC')
    kw.setdefault('peso_toneladas', Decimal('9.82'))
    kw.setdefault('contenedor', 'TLLU7742369')
    kw.setdefault('fecha_modulacion_aduana', FECHA)
    return Modulacion.objects.create(**kw)


def _rows(ws):
    return list(ws.iter_rows(values_only=True))


class ConstruirProgramaDespachoTests(TestCase):
    def test_encabezados_y_registro_combinado(self):
        wb = construir_programa_despacho(FECHA)
        ws = wb.active
        self.assertEqual(ws.title, 'Programa de despacho')
        fila1 = _rows(ws)[0]
        self.assertEqual(fila1[0], 'FECHA DE DESPACHO')
        self.assertEqual(fila1[6], 'OPERADOR')
        self.assertEqual(fila1[7], 'CLIENTE')
        self.assertEqual(fila1[12], 'SELLO')
        self.assertEqual(fila1[13], 'Nº DE MANIOBRA')
        # J1:L1 combinado = "REGISTRO"
        self.assertEqual(ws['J1'].value, 'REGISTRO')

    def test_fila_de_datos_mapea_los_campos(self):
        cli = Cliente.objects.create(nombre='Transportes Moya', alias='MOY')
        op = Operador.objects.create(nombre='Jose Alberto Garcia', tipo='LOCAL', activo=True)
        uni = Unidad.objects.create(
            numero_economico='LE08', placa='ND0814D', tipo='LOCAL', activa=True,
            año=2020, capacidad_combustible=Decimal('300'),
            rendimiento_esperado=Decimal('2.5'),
        )
        _mod(cliente=cli, operador=op, unidad=uni, carril='5C', sello_colocado=False,
             hora_registro=_aware(2026, 8, 28, 9, 30),
             hora_ingreso=_aware(2026, 8, 28, 11, 0),
             hora_carga=_aware(2026, 8, 28, 12, 30))
        ws = construir_programa_despacho(FECHA).active
        datos = [r for r in _rows(ws) if r[5] == 'TLLU7742369'][0]
        self.assertEqual(datos[0], '28-ago-26')
        self.assertEqual(datos[1], 'LCTPC')
        self.assertEqual(datos[2], 'LOGINCO')
        self.assertEqual(datos[3], '40')
        self.assertEqual(datos[4], 9.82)
        self.assertIn('Jose Alberto Garcia', datos[6])
        self.assertIn('ECO LE08 PLACAS ND0814D', datos[6])
        self.assertEqual(datos[7], 'MOY')
        self.assertEqual(datos[8], '5C')
        self.assertEqual(datos[12], 'NO')
        self.assertEqual(datos[13], 1)

    def test_tipo_toma_solo_los_digitos_iniciales(self):
        _mod(tipo_contenedor='45G1', contenedor='ZZZU9999999')
        ws = construir_programa_despacho(FECHA).active
        datos = [r for r in _rows(ws) if r[5] == 'ZZZU9999999'][0]
        self.assertEqual(datos[3], '45')

    def test_tres_horas_llevan_fondo(self):
        _mod(hora_registro=_aware(2026, 8, 28, 9, 30),
             hora_ingreso=_aware(2026, 8, 28, 11, 0),
             hora_carga=_aware(2026, 8, 28, 12, 30))
        ws = construir_programa_despacho(FECHA).active
        fila = [i for i, r in enumerate(_rows(ws), start=1) if r[5] == 'TLLU7742369'][0]
        self.assertIn('9:30 a. m.', ws.cell(row=fila, column=10).value)
        self.assertIn('11:00 a. m.', ws.cell(row=fila, column=11).value)
        self.assertIn('12:30 p. m.', ws.cell(row=fila, column=12).value)
        self.assertEqual(ws.cell(row=fila, column=10).fill.fill_type, 'solid')
        self.assertEqual(ws.cell(row=fila, column=12).fill.fill_type, 'solid')
        # (HOY) porque la fecha de la hora == fecha del reporte
        self.assertIn('(HOY)', ws.cell(row=fila, column=10).value)

    def test_una_sola_hora_se_combina_sin_fondo(self):
        _mod(hora_ingreso=_aware(2026, 8, 28, 2, 0))
        ws = construir_programa_despacho(FECHA).active
        fila = [i for i, r in enumerate(_rows(ws), start=1) if r[5] == 'TLLU7742369'][0]
        self.assertIn('2:00 a. m.', ws.cell(row=fila, column=10).value)
        self.assertIn(ws.cell(row=fila, column=10).fill.fill_type, (None, 'none'))

    def test_sin_horas_dice_cita_pendiente(self):
        _mod()
        ws = construir_programa_despacho(FECHA).active
        fila = [i for i, r in enumerate(_rows(ws), start=1) if r[5] == 'TLLU7742369'][0]
        self.assertEqual(ws.cell(row=fila, column=10).value, 'CITA PENDIENTE')

    def test_agrupa_por_alias_con_fila_encabezado(self):
        # 'Zeta' (sin alias) empieza con Z pero por orden alfabético normal iría
        # antes de 'sin cliente'; el bloque sin cliente debe ir SIEMPRE al final.
        moy = Cliente.objects.create(nombre='Moya', alias='MOY')
        nol = Cliente.objects.create(nombre='Nolasco SA')  # sin alias
        zeta = Cliente.objects.create(nombre='Zeta SA')  # sin alias
        _mod(cliente=moy, contenedor='AAAU1111111')
        _mod(cliente=moy, contenedor='AAAU2222222')
        _mod(cliente=nol, contenedor='BBBU3333333')
        _mod(cliente=zeta, contenedor='CCCU4444444')
        _mod(cliente=None, contenedor='DDDU5555555')
        ws = construir_programa_despacho(FECHA).active
        filas = _rows(ws)
        # las filas-encabezado salen en el orden esperado: alias, luego el
        # bloque "sin cliente" al final, sin importar el orden alfabético.
        encabezados_grupo = [r[0] for r in filas if r[0] and 'maniobra(s)' in str(r[0])]
        self.assertEqual(encabezados_grupo, [
            'MOY — 2 maniobra(s)',
            'Nolasco SA — 1 maniobra(s)',
            'Zeta SA — 1 maniobra(s)',
            '— — 1 maniobra(s)',
        ])
        # Nº de maniobra corre 1..5 en el orden de las filas (no cuenta los
        # encabezados de grupo) y el de "sin cliente" es el último.
        contenedores = [
            'AAAU1111111', 'AAAU2222222', 'BBBU3333333', 'CCCU4444444', 'DDDU5555555',
        ]
        nums = [r[13] for r in filas if r[5] in contenedores]
        self.assertEqual(nums, [1, 2, 3, 4, 5])
        fila_sin_cliente = [r for r in filas if r[5] == 'DDDU5555555'][0]
        self.assertEqual(fila_sin_cliente[7], '—')

    def test_sin_operador_ni_unidad_no_revienta(self):
        _mod()  # sin operador ni unidad
        ws = construir_programa_despacho(FECHA).active
        datos = [r for r in _rows(ws) if r[5] == 'TLLU7742369'][0]
        self.assertEqual(datos[6], '')

    def test_peso_es_numero(self):
        _mod(peso_toneladas=Decimal('10.37'))
        ws = construir_programa_despacho(FECHA).active
        fila = [i for i, r in enumerate(_rows(ws), start=1) if r[5] == 'TLLU7742369'][0]
        c = ws.cell(row=fila, column=5)
        self.assertEqual(c.value, 10.37)
        self.assertEqual(c.number_format, '0.00')


class ReporteDespachoViewTests(TestCase):
    def setUp(self):
        self.user = get_user_model().objects.create_user('u', 'u@e.com', 'pw')

    def test_form_requiere_login(self):
        resp = self.client.get(reverse('modulacion:reporte_despacho'))
        self.assertEqual(resp.status_code, 302)

    def test_form_muestra_input_de_fecha_con_hoy(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse('modulacion:reporte_despacho'))
        self.assertEqual(resp.status_code, 200)
        hoy = timezone.localdate().isoformat()
        self.assertContains(resp, 'type="date"')
        self.assertContains(resp, f'value="{hoy}"')

    def test_form_usa_fecha_de_querystring(self):
        # El input de fecha (compartido por los botones Excel y WhatsApp,
        # al vivir en un solo <form>) debe reflejar el ?fecha= de la URL,
        # no siempre "hoy".
        self.client.force_login(self.user)
        resp = self.client.get(reverse('modulacion:reporte_despacho'), {'fecha': FECHA.isoformat()})
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, f'value="{FECHA.isoformat()}"')
        self.assertNotContains(resp, f'value="{timezone.localdate().isoformat()}"')

    def test_form_fecha_invalida_en_querystring_usa_hoy(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse('modulacion:reporte_despacho'), {'fecha': 'no-es-fecha'})
        self.assertEqual(resp.status_code, 200)
        hoy = timezone.localdate().isoformat()
        self.assertContains(resp, f'value="{hoy}"')

    def test_descarga_xlsx(self):
        self.client.force_login(self.user)
        _mod()
        resp = self.client.get(
            reverse('modulacion:reporte_despacho_xlsx'), {'fecha': '2026-08-28'})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        self.assertIn('attachment; filename="programa_despacho_2026-08-28.xlsx"',
                      resp['Content-Disposition'])
        wb = openpyxl.load_workbook(BytesIO(resp.getvalue()))
        self.assertEqual(wb.active['A1'].value, 'FECHA DE DESPACHO')

    def test_descarga_sin_fecha_usa_hoy(self):
        self.client.force_login(self.user)
        resp = self.client.get(reverse('modulacion:reporte_despacho_xlsx'))
        self.assertEqual(resp.status_code, 200)
        hoy = timezone.localdate().isoformat()
        self.assertIn(f'programa_despacho_{hoy}.xlsx', resp['Content-Disposition'])

    def test_descarga_fecha_invalida_usa_hoy(self):
        self.client.force_login(self.user)
        resp = self.client.get(
            reverse('modulacion:reporte_despacho_xlsx'), {'fecha': 'no-es-fecha'})
        self.assertEqual(resp.status_code, 200)


from modulos.modulacion.reportes import _agrupar_y_numerar


class AgruparYNumerarTests(TestCase):
    def test_agrupa_y_numera_igual_que_el_xlsx(self):
        moy = Cliente.objects.create(nombre='Moya', alias='MOY')
        nol = Cliente.objects.create(nombre='Nolasco SA')  # sin alias
        _mod(cliente=moy, contenedor='AAAU1111111', hora_registro=_aware(2026, 8, 28, 9, 0))
        _mod(cliente=moy, contenedor='AAAU2222222', hora_registro=_aware(2026, 8, 28, 10, 0))
        _mod(cliente=nol, contenedor='BBBU3333333')
        _mod(cliente=None, contenedor='CCCU4444444')

        grupos = _agrupar_y_numerar(FECHA)

        etiquetas = [g[0] for g in grupos]
        self.assertEqual(etiquetas, ['MOY', 'Nolasco SA', '—'])

        cliente_ids = [g[1] for g in grupos]
        self.assertEqual(cliente_ids, [moy.pk, nol.pk, None])

        # numeración corrida 1..4 sobre todo el reporte, no por grupo
        todos_los_numeros = [num for _e, _c, items in grupos for num, _m in items]
        self.assertEqual(todos_los_numeros, [1, 2, 3, 4])

        contenedores_moy = [m.contenedor for _num, m in grupos[0][2]]
        self.assertEqual(contenedores_moy, ['AAAU1111111', 'AAAU2222222'])
